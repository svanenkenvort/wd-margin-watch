#!/usr/bin/env python3
"""Import the MASTER TRACKER workbook into data/store.json (one-time seed).

Usage: python3 scripts/import_xlsx.py "MASTER TRACKER 2025 Template 1.xlsx"

Reads the 'Open Jobs', '2026 Completed Jobs' and '2025 Completed Jobs' tabs.
Only the raw inputs are imported; every calculated column is recomputed by the app.
The production stage is read from the colour of the job-name cell (your legend on the Summary tab).
"""
import sys, json, uuid, datetime as dt, re
import openpyxl

SRC = sys.argv[1]
OUT = sys.argv[2] if len(sys.argv) > 2 else "data/store.json"

STAGE_BY_RGB = {
    "FFFFFF00": "ordered",
    "FFFF0000": "toOrder",
    "FF92D050": "toSchedule",
    "FF00B0F0": "scheduled",
    "FFFFC000": "completedAwaitingPay",
    "FF7030A0": "needsMeasure",
}
GREY_AWAITING_DEPOSIT = -0.25
GREY_AWAITING_FINAL = -0.35

def stage_of(cell):
    f = cell.fill
    if not f or not f.fill_type:
        return None
    c = f.fgColor
    if c.type == "rgb":
        return STAGE_BY_RGB.get(c.rgb)
    if c.type == "theme" and c.theme == 0:
        t = round(c.tint or 0, 2)
        if t == GREY_AWAITING_DEPOSIT:
            return "awaitingDeposit"
        if t == GREY_AWAITING_FINAL:
            return "awaitingFinal"
    return None

def num(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace("$", "").replace(",", "").strip())
    except ValueError:
        return None

def iso(v):
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    return None

def clean_product(v):
    s = str(v or "").upper().replace("'S", "").strip()
    s = re.sub(r"\s*/\s*", "/", s)
    s = re.sub(r"\s+", " ", s)
    s = {"PD W": "W/PD", "SELF INSTALL": "SELF"}.get(s, s)
    return s

def clean_rep(v):
    s = re.sub(r"\s+", " ", str(v or "").strip().upper())
    return "SHAUN V" if s == "SHAUN" else s

wb_v = openpyxl.load_workbook(SRC, data_only=True)      # cached values
wb_f = openpyxl.load_workbook(SRC, data_only=False)     # fills + formulas

jobs, overrides = [], 0

def read_tab(name, state, year):
    global overrides
    wv, wf = wb_v[name], wb_f[name]
    for r in range(6, wv.max_row + 1):
        jn = wv.cell(r, 1).value
        g = num(wv.cell(r, 7).value)
        if jn in (None, "") or g is None:
            continue
        col = lambda c: wv.cell(r, c).value          # noqa: E731
        dep_date_raw = col(11)
        financed = isinstance(dep_date_raw, str) and "FINANC" in dep_date_raw.upper()
        T, U, Y = num(col(20)), num(col(21)), num(col(25))
        job = {
            "id": str(uuid.uuid4()),
            "jobNumber": str(jn).strip(),
            "name": str(col(3) or "").strip(),
            "product": clean_product(col(2)),
            "rep": clean_rep(col(28)),
            "soldDate": iso(col(4)),
            "state": state,
            "stage": stage_of(wf.cell(r, 3)) if state == "open" else None,
            "completeDate": iso(col(6)) if state == "completed" else None,
            "reportYear": year if state == "completed" else None,
            "contractValue": g,
            "financePct": num(col(8)) or 0.0,
            "financed": financed,
            "depositAmount": num(col(10)) or 0.0,
            "depositDate": iso(dep_date_raw),
            "stageAmount": num(col(14)) or 0.0,
            "stageDate": iso(col(15)),
            "finalAmount": num(col(16)) or 0.0,
            "finalDate": iso(col(17)),
            "materialsPaid": num(col(12)),
            "materialsEst": T,
            "materialsActual": None,
            "laborEst": U,
            "laborActual": Y,
            "commissionRate": num(col(29)),
            "commissionPaid": (num(col(31)) or 0.0) + (num(col(33)) or 0.0),
            "commissionPaidDate": iso(col(34)) or iso(col(32)),
            "notes": str(col(38) or "").strip() if state == "open" else str(col(38) or "").strip(),
            "source": "sheet",
            "builderPrimeId": None,
            "updatedAt": dt.datetime.now(dt.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
        }
        x = num(col(24))
        if x is not None and abs(x - (0.03 * g + (T or 0.0))) > 0.01:
            overrides += 1
            job["materialsActual"] = x   # typed over the formula in the sheet
        jobs.append(job)

read_tab("Open Jobs", "open", None)
read_tab("2026 Completed Jobs", "completed", 2026)
read_tab("2025 Completed Jobs", "completed", 2025)

store = {
    "version": 1,
    "settings": {
        "targetMargin": 0.4,
        "marginBasis": "afterCommission",
        "materialsAllowance": 0.03,
        "reps": [
            {"name": "SCHMITZ", "commissionRate": 0.10},
            {"name": "SHAUN V", "commissionRate": 0.0},
        ],
        "builderPrime": {"subdomain": "windowdepotnew", "lastSync": None},
    },
    "jobs": jobs,
}
with open(OUT, "w") as fh:
    json.dump(store, fh, indent=1)

from collections import Counter
print("jobs:", len(jobs), dict(Counter((j["state"], j["reportYear"]) for j in jobs)))
print("stages:", dict(Counter(j["stage"] for j in jobs if j["state"] == "open")))
print("reps:", dict(Counter(j["rep"] for j in jobs)))
print("products:", dict(Counter(j["product"] for j in jobs).most_common(20)))
print("materials-actual overrides (X != 3% + T):", overrides)
print("financed deposits:", sum(j["financed"] for j in jobs))
